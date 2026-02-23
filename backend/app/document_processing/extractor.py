"""
Multi-format document text extractor.

Supports: PDF, DOCX, PPTX, XLSX, CSV, and images (with OCR).
Uses the 'unstructured' library for structured extraction with
special handling for tables to preserve row/column structure.
"""
import asyncio
import csv
import io
import logging
import os
from enum import Enum
from html.parser import HTMLParser
from pathlib import Path
from dataclasses import dataclass, field

logger = logging.getLogger(__name__)


class ElementType(str, Enum):
    """Types of content elements extracted from documents."""
    TEXT = "text"
    TABLE = "table"
    TITLE = "title"
    LIST_ITEM = "list_item"


@dataclass
class ContentElement:
    """A typed, structured content element from a document."""
    element_type: ElementType
    text: str
    metadata: dict = field(default_factory=dict)


@dataclass
class ExtractedPage:
    """A single page/sheet/slide of extracted content."""
    page_number: int
    text: str
    elements: list[ContentElement] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)


@dataclass
class ExtractedDocument:
    """Result of document extraction."""
    filename: str
    file_type: str
    pages: list[ExtractedPage] = field(default_factory=list)
    total_text: str = ""
    metadata: dict = field(default_factory=dict)
    error: str | None = None

    @property
    def page_count(self) -> int:
        return len(self.pages)

    def full_text(self) -> str:
        """Get all text combined."""
        if self.total_text:
            return self.total_text
        return "\n\n".join(p.text for p in self.pages if p.text)


class _HTMLTableParser(HTMLParser):
    """Parse HTML <table> into a list of rows (each row is a list of cell strings)."""

    def __init__(self):
        super().__init__()
        self.rows: list[list[str]] = []
        self._current_row: list[str] = []
        self._current_cell: list[str] = []
        self._in_cell = False

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._current_row = []
        elif tag in ("td", "th"):
            self._current_cell = []
            self._in_cell = True

    def handle_endtag(self, tag):
        if tag in ("td", "th"):
            self._in_cell = False
            self._current_row.append("".join(self._current_cell).strip())
        elif tag == "tr":
            if self._current_row:
                self.rows.append(self._current_row)

    def handle_data(self, data):
        if self._in_cell:
            self._current_cell.append(data)


def _html_table_to_markdown(html: str) -> str:
    """Convert an HTML <table> string to a markdown pipe-delimited table."""
    parser = _HTMLTableParser()
    parser.feed(html)
    rows = parser.rows

    if not rows:
        return ""

    # Determine max columns for alignment
    max_cols = max(len(r) for r in rows)

    # Normalize all rows to same column count
    normalized = []
    for row in rows:
        padded = row + [""] * (max_cols - len(row))
        normalized.append(padded)

    lines = []
    # Header row
    header = normalized[0]
    lines.append("| " + " | ".join(header) + " |")
    lines.append("| " + " | ".join("---" for _ in header) + " |")

    # Data rows
    for row in normalized[1:]:
        lines.append("| " + " | ".join(row) + " |")

    return "\n".join(lines)


def _rows_to_markdown(rows: list[list[str]]) -> str:
    """Convert a list of rows to a markdown pipe-delimited table."""
    if not rows:
        return ""

    max_cols = max(len(r) for r in rows)
    normalized = []
    for row in rows:
        padded = [str(cell) if cell is not None else "" for cell in row] + [""] * (max_cols - len(row))
        normalized.append(padded)

    lines = []
    header = normalized[0]
    lines.append("| " + " | ".join(header) + " |")
    lines.append("| " + " | ".join("---" for _ in header) + " |")

    for row in normalized[1:]:
        lines.append("| " + " | ".join(row) + " |")

    return "\n".join(lines)


class DocumentExtractor:
    """
    Extract text from multiple document formats.

    Usage:
        extractor = DocumentExtractor()
        result = extractor.extract("/path/to/doc.pdf", "pdf")

        # With vision-enhanced table extraction for PDFs:
        extractor = DocumentExtractor(
            use_vision_tables=True,
            vision_backend=get_vision_backend(),
        )
    """

    SUPPORTED_TYPES = {"pdf", "docx", "pptx", "xlsx", "csv", "image", "other"}

    def __init__(self, use_vision_tables: bool = False, vision_backend=None):
        self._use_vision_tables = use_vision_tables
        self._vision_backend = vision_backend

    def extract(self, file_path: str, file_type: str) -> ExtractedDocument:
        """
        Extract text from a document.

        Args:
            file_path: Absolute path to the file
            file_type: One of: pdf, docx, pptx, xlsx, csv, image, other

        Returns:
            ExtractedDocument with extracted text and metadata
        """
        path = Path(file_path)
        if not path.exists():
            return ExtractedDocument(
                filename=path.name,
                file_type=file_type,
                error=f"File not found: {file_path}",
            )

        if file_type not in self.SUPPORTED_TYPES:
            return ExtractedDocument(
                filename=path.name,
                file_type=file_type,
                error=f"Unsupported file type: {file_type}",
            )

        try:
            method = getattr(self, f"_extract_{file_type}")
            return method(path)
        except Exception as e:
            logger.exception(f"Error extracting {file_path}")
            return ExtractedDocument(
                filename=path.name,
                file_type=file_type,
                error=str(e),
            )

    def _extract_pdf(self, path: Path) -> ExtractedDocument:
        """Extract text from PDF using unstructured with hi_res for table detection."""
        if self._use_vision_tables and self._vision_backend:
            return self._extract_pdf_with_vision(path)

        from unstructured.partition.pdf import partition_pdf

        elements = partition_pdf(
            str(path),
            strategy="hi_res",
            infer_table_structure=True,
        )
        return self._elements_to_document(path, "pdf", elements)

    def _extract_pdf_with_vision(self, path: Path) -> ExtractedDocument:
        """
        Simplified vision PDF extraction — per-page routing.

        Each page is independently routed:
        - Pages with images/visual content → VLM full-page extraction
        - Pages with only digital text → PyMuPDF text extraction
        """
        from app.document_processing.page_renderer import (
            analyze_pages, render_page_to_image,
        )

        page_analysis = analyze_pages(str(path))
        vlm_pages = [p for p in page_analysis if p["method"] == "vlm"]
        text_pages = [p for p in page_analysis if p["method"] == "text"]

        logger.info(
            "PDF %s: %d pages — %d VLM, %d text-only",
            path.name, len(page_analysis), len(vlm_pages), len(text_pages),
        )

        extracted_pages = []
        temp_files = []

        try:
            loop = asyncio.new_event_loop()
            try:
                for page_info in page_analysis:
                    page_num = page_info["page_num"]

                    if page_info["method"] == "vlm":
                        # Render page and send to VLM
                        page_img = render_page_to_image(str(path), page_num, dpi=300)
                        temp_files.append(page_img)

                        try:
                            page_text = loop.run_until_complete(
                                self._vision_backend.extract_page(page_img)
                            )
                        except Exception:
                            logger.exception(
                                "VLM failed on page %d of %s", page_num, path.name,
                            )
                            page_text = ""

                        elements = self._parse_vlm_page_output(page_text)
                        pipeline = "vlm"

                    else:
                        # Pure text page — extract with PyMuPDF
                        page_text, elements = self._extract_text_page(
                            str(path), page_num,
                        )
                        pipeline = "text"

                    extracted_pages.append(ExtractedPage(
                        page_number=page_num,
                        text=page_text,
                        elements=elements,
                        metadata={
                            "source_file": path.name,
                            "pipeline": pipeline,
                            "has_images": page_info["has_images"],
                            "has_tables": page_info["has_tables"],
                        },
                    ))

                    logger.info(
                        "Page %d/%d [%s]: %d chars, %d elements",
                        page_num, len(page_analysis), pipeline,
                        len(page_text), len(elements),
                    )
            finally:
                loop.close()

            doc_result = ExtractedDocument(
                filename=path.name,
                file_type="pdf",
                pages=extracted_pages,
                metadata={
                    "source_path": str(path),
                    "page_count": len(extracted_pages),
                    "pipeline": "per_page_routing",
                    "vlm_pages": [p["page_num"] for p in vlm_pages],
                    "text_pages": [p["page_num"] for p in text_pages],
                },
            )
            doc_result.total_text = doc_result.full_text()
            return doc_result

        finally:
            for tmp in temp_files:
                try:
                    os.unlink(tmp)
                except OSError:
                    pass

    def _extract_text_page(
        self, pdf_path: str, page_num: int,
    ) -> tuple[str, list[ContentElement]]:
        """Extract text from a single digital page using PyMuPDF."""
        import fitz

        doc = fitz.open(pdf_path)
        try:
            page = doc[page_num - 1]  # 0-indexed
            blocks = page.get_text("dict")["blocks"]
            elements = []
            text_parts = []

            for block in blocks:
                if block["type"] != 0:  # only text blocks
                    continue

                block_text = ""
                for line in block.get("lines", []):
                    for span in line.get("spans", []):
                        block_text += span.get("text", "")
                    block_text += "\n"
                block_text = block_text.strip()

                if not block_text:
                    continue

                # Heuristic: short bold text → title
                spans = [
                    s for l in block.get("lines", []) for s in l.get("spans", [])
                ]
                is_title = (
                    len(block_text) < 200
                    and spans
                    and any(s.get("flags", 0) & 2**4 for s in spans)  # bold flag
                )
                el_type = ElementType.TITLE if is_title else ElementType.TEXT
                elements.append(ContentElement(element_type=el_type, text=block_text))

                if el_type == ElementType.TITLE:
                    text_parts.append(f"\n## {block_text}\n")
                else:
                    text_parts.append(block_text)

            return "\n".join(text_parts), elements
        finally:
            doc.close()

    @staticmethod
    def _parse_vlm_page_output(text: str) -> list[ContentElement]:
        """
        Parse VLM full-page output into typed ContentElement instances.

        - Lines starting with | → TABLE elements (grouped into contiguous blocks).
        - Lines starting with # → TITLE elements.
        - Everything else → TEXT elements (grouped into contiguous blocks).
        """
        if not text or not text.strip():
            return []

        elements: list[ContentElement] = []
        lines = text.split("\n")

        current_type: ElementType | None = None
        current_lines: list[str] = []

        def _flush():
            nonlocal current_type, current_lines
            if current_type is None or not current_lines:
                current_lines = []
                return
            block = "\n".join(current_lines).strip()
            if block:
                elements.append(ContentElement(
                    element_type=current_type,
                    text=block,
                ))
            current_lines = []
            current_type = None

        for line in lines:
            stripped = line.strip()

            if stripped.startswith("|"):
                # Table line (content or separator like |---|---|)
                if current_type != ElementType.TABLE:
                    _flush()
                    current_type = ElementType.TABLE
                current_lines.append(line)

            elif stripped.startswith("#"):
                # Heading — each heading is its own element
                _flush()
                elements.append(ContentElement(
                    element_type=ElementType.TITLE,
                    text=stripped.lstrip("#").strip(),
                ))

            else:
                # Regular text
                if current_type != ElementType.TEXT:
                    _flush()
                    current_type = ElementType.TEXT
                current_lines.append(line)

        _flush()
        return elements

    def _extract_docx(self, path: Path) -> ExtractedDocument:
        """Extract text from Word documents."""
        from unstructured.partition.docx import partition_docx

        elements = partition_docx(str(path))
        return self._elements_to_document(path, "docx", elements)

    def _extract_pptx(self, path: Path) -> ExtractedDocument:
        """Extract text from PowerPoint."""
        from unstructured.partition.pptx import partition_pptx

        elements = partition_pptx(str(path))
        return self._elements_to_document(path, "pptx", elements)

    # Max rows per table chunk for Excel/CSV — keeps chunks manageable
    # for embedding and retrieval. ~50 rows × ~80 chars = ~4000 chars.
    XLSX_ROWS_PER_CHUNK = 50

    def _extract_xlsx(self, path: Path) -> ExtractedDocument:
        """
        Extract text from Excel spreadsheets using openpyxl.

        Large sheets are pre-chunked at extraction time: each chunk contains
        the header row + up to XLSX_ROWS_PER_CHUNK data rows. This avoids
        building a single giant markdown string and produces a reasonable
        number of chunks even for sheets with 50k+ rows.
        """
        import openpyxl

        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        pages = []
        total_rows = 0

        for sheet_idx, sheet_name in enumerate(wb.sheetnames, start=1):
            ws = wb[sheet_name]

            # Collect all non-empty rows, stringify cells
            all_rows: list[list[str]] = []
            for row in ws.iter_rows(values_only=True):
                if any(cell is not None for cell in row):
                    all_rows.append([str(cell) if cell is not None else "" for cell in row])

            if not all_rows:
                continue

            total_rows += len(all_rows)

            # First row is always the header
            header_row = all_rows[0]
            data_rows = all_rows[1:]

            # Pre-chunk: split data rows into groups of XLSX_ROWS_PER_CHUNK
            elements: list[ContentElement] = []

            if not data_rows:
                # Header-only sheet
                md = _rows_to_markdown([header_row])
                elements.append(ContentElement(
                    element_type=ElementType.TABLE,
                    text=md,
                    metadata={
                        "sheet_name": sheet_name,
                        "section_header": f"Sheet: {sheet_name}",
                        "row_range": "header only",
                    },
                ))
            else:
                for start in range(0, len(data_rows), self.XLSX_ROWS_PER_CHUNK):
                    batch = data_rows[start:start + self.XLSX_ROWS_PER_CHUNK]
                    # Always include header row so each chunk is self-contained
                    chunk_rows = [header_row] + batch
                    md = _rows_to_markdown(chunk_rows)

                    row_start = start + 2  # +2 because row 1 is header (1-indexed)
                    row_end = row_start + len(batch) - 1

                    elements.append(ContentElement(
                        element_type=ElementType.TABLE,
                        text=md,
                        metadata={
                            "sheet_name": sheet_name,
                            "section_header": f"Sheet: {sheet_name}",
                            "row_range": f"rows {row_start}-{row_end}",
                        },
                    ))

            # Build page text for full_text() (first chunk only for summary)
            page_text = f"### Sheet: {sheet_name} ({len(all_rows)} rows)\n\n{elements[0].text}"
            if len(elements) > 1:
                page_text += f"\n\n... and {len(elements) - 1} more table segments"

            pages.append(ExtractedPage(
                page_number=sheet_idx,
                text=page_text,
                elements=elements,
                metadata={"source_file": path.name, "sheet_name": sheet_name},
            ))

        wb.close()

        doc = ExtractedDocument(
            filename=path.name,
            file_type="xlsx",
            pages=pages,
            metadata={
                "source_path": str(path),
                "page_count": len(pages),
                "total_rows": total_rows,
            },
        )
        doc.total_text = doc.full_text()

        logger.info(
            f"Extracted {path.name}: {len(pages)} sheets, {total_rows} total rows, "
            f"{sum(len(p.elements) for p in pages)} table segments"
        )
        return doc

    def _extract_csv(self, path: Path) -> ExtractedDocument:
        """Extract text from CSV files, pre-chunked for large files."""
        text = path.read_text(encoding="utf-8-sig")
        reader = csv.reader(io.StringIO(text))
        all_rows = [row for row in reader if any(cell.strip() for cell in row)]

        if not all_rows:
            return ExtractedDocument(
                filename=path.name,
                file_type="csv",
                error="CSV file is empty",
            )

        header_row = all_rows[0]
        data_rows = all_rows[1:]

        elements: list[ContentElement] = []

        if not data_rows:
            md = _rows_to_markdown([header_row])
            elements.append(ContentElement(
                element_type=ElementType.TABLE,
                text=md,
                metadata={"row_count": 1},
            ))
        else:
            for start in range(0, len(data_rows), self.XLSX_ROWS_PER_CHUNK):
                batch = data_rows[start:start + self.XLSX_ROWS_PER_CHUNK]
                chunk_rows = [header_row] + batch
                md = _rows_to_markdown(chunk_rows)

                row_start = start + 2
                row_end = row_start + len(batch) - 1

                elements.append(ContentElement(
                    element_type=ElementType.TABLE,
                    text=md,
                    metadata={
                        "row_count": len(batch),
                        "row_range": f"rows {row_start}-{row_end}",
                    },
                ))

        page = ExtractedPage(
            page_number=1,
            text=elements[0].text,
            elements=elements,
            metadata={"source_file": path.name},
        )

        doc = ExtractedDocument(
            filename=path.name,
            file_type="csv",
            pages=[page],
            metadata={"source_path": str(path), "page_count": 1, "total_rows": len(all_rows)},
        )
        doc.total_text = doc.full_text()

        logger.info(
            f"Extracted {path.name}: {len(all_rows)} rows, "
            f"{len(elements)} table segments"
        )
        return doc

    def _extract_image(self, path: Path) -> ExtractedDocument:
        """Extract text from images using OCR."""
        from unstructured.partition.image import partition_image

        elements = partition_image(str(path), strategy="ocr_only")
        return self._elements_to_document(path, "image", elements)

    def _extract_other(self, path: Path) -> ExtractedDocument:
        """Extract text from plain text files."""
        try:
            text = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            text = path.read_text(encoding="latin-1")

        page = ExtractedPage(
            page_number=1,
            text=text,
            elements=[ContentElement(element_type=ElementType.TEXT, text=text)],
            metadata={"source_file": path.name},
        )

        doc = ExtractedDocument(
            filename=path.name,
            file_type="other",
            pages=[page],
            total_text=text,
            metadata={"source_path": str(path), "page_count": 1},
        )
        return doc

    def _classify_element(self, element) -> tuple[ElementType, str]:
        """Classify an unstructured element and extract its text with structure."""
        element_class = type(element).__name__

        if element_class in ("Table", "TableChunk"):
            # Prefer vision-extracted table if available
            vision_text = getattr(element, "_vision_table_text", None)
            if vision_text:
                return ElementType.TABLE, vision_text

            # Try to get HTML representation for accurate table conversion
            html = getattr(element.metadata, "text_as_html", None)
            if html:
                md_table = _html_table_to_markdown(html)
                if md_table:
                    return ElementType.TABLE, md_table
            # Fallback: wrap in fenced code block to preserve spacing
            return ElementType.TABLE, f"```\n{str(element)}\n```"

        if element_class == "Image":
            # Use vision caption if available
            vision_caption = getattr(element, "_vision_caption", None)
            if vision_caption:
                return ElementType.TEXT, vision_caption

        if element_class == "Title":
            return ElementType.TITLE, str(element)

        if element_class == "ListItem":
            return ElementType.LIST_ITEM, f"- {str(element)}"

        # NarrativeText, UncategorizedText, etc.
        return ElementType.TEXT, str(element)

    def _elements_to_document(
        self, path: Path, file_type: str, elements: list
    ) -> ExtractedDocument:
        """Convert unstructured elements to our ExtractedDocument format with structure."""
        pages: dict[int, list[ContentElement]] = {}
        current_section: str | None = None

        for element in elements:
            page_num = getattr(element.metadata, "page_number", 1) or 1
            if page_num not in pages:
                pages[page_num] = []

            el_type, text = self._classify_element(element)

            # Track section headers
            if el_type == ElementType.TITLE:
                current_section = text

            content_el = ContentElement(
                element_type=el_type,
                text=text,
                metadata={"section_header": current_section},
            )
            pages[page_num].append(content_el)

        extracted_pages = []
        for num, elements_list in sorted(pages.items()):
            # Build plain text from elements
            text_parts = []
            for el in elements_list:
                if el.element_type == ElementType.TITLE:
                    text_parts.append(f"\n## {el.text}\n")
                else:
                    text_parts.append(el.text)

            page_text = "\n".join(text_parts)

            extracted_pages.append(ExtractedPage(
                page_number=num,
                text=page_text,
                elements=elements_list,
                metadata={"source_file": path.name},
            ))

        doc = ExtractedDocument(
            filename=path.name,
            file_type=file_type,
            pages=extracted_pages,
            metadata={
                "source_path": str(path),
                "page_count": len(extracted_pages),
            },
        )
        doc.total_text = doc.full_text()

        logger.info(
            f"Extracted {path.name}: {len(extracted_pages)} pages, "
            f"{len(doc.total_text)} chars"
        )
        return doc
